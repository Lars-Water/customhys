from collections import Counter
import pandas as pd
import os

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule


def _get_excel_column_name(idx):
    """
    Convert a 0-based index to an Excel column name (e.g., 0 -> 'A', 25 -> 'Z', 26 -> 'AA', etc.).

    Parameters:
        idx (int): 0-based index of the column.

    Returns:
        str: Excel column name.
    """
    alphabet = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
    result = []
    while idx >= 0:
        idx, remainder = divmod(idx, 26)
        result.append(alphabet[remainder])
        idx -= 1
    return ''.join(reversed(result))


def format_configuration_sheet(file_path, positions_df, nr_of_backbone_switches):
    # Load the workbook and select the active worksheet
    wb = load_workbook(file_path)
    if wb is None:
        raise ValueError("Failed to load workbook.")
    ws = wb.active
    if ws is None:
        raise ValueError("Failed to access the active worksheet.")

    # Define conditional formatting rules
    light_red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
    light_blue_fill = PatternFill(start_color='CCCCFF', end_color='CCCCFF', fill_type='solid')
    light_green_fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
    light_yellow_fill = PatternFill(start_color='FFFFCC', end_color='FFFFCC', fill_type='solid')

    # List of conditional formatting rules for each value
    rules = [
        (1, light_red_fill),
        (2, light_blue_fill),
        (3, light_green_fill),
        (4, light_yellow_fill)
    ]

    # Apply the conditional formatting rules to each "Position" column
    for i in range(1, nr_of_backbone_switches):  # Starting from 1 because the first column is 'Category'
        col_letter = _get_excel_column_name(i)
        for value, fill in rules:
            rule = CellIsRule(operator='equal', formula=[str(value)], fill=fill)
            range_string = f'{col_letter}2:{col_letter}{len(positions_df) + 1}'  # Construct range string dynamically
            ws.conditional_formatting.add(range_string, rule)

    # Save the workbook
    wb.save(file_path)

    print(f"Conditional formatting applied and saved to {file_path}")


def write_processed_design_points_to_xlsx(hh_run_path, optimal_configurations_list, nr_of_backbone_switches):
    # Convert to a DataFrame
    columns = ['Category'] + [f'Position_{i+1}' for i in range(nr_of_backbone_switches-1)]
    positions_df = pd.DataFrame(optimal_configurations_list, columns=columns)

    # Identify all unique cable values accross all configuration points.
    unique_values = set()
    for row in positions_df.itertuples(index=False):
        unique_values.update(row[1:])  # Skip the 'Category' column

    # Calculate occurrences for each row, ensuring all unique values are included
    counts_list = []
    for row in positions_df.itertuples(index=False):
        count_dict = Counter(row[1:])  # Exclude 'Category' column from count
        count_dict_full = {val: count_dict.get(val, 0) for val in sorted(unique_values)}  # Include zero counts
        counts_list.append([row[0]] + list(count_dict_full.values()))

    # Convert the counts list to a DataFrame
    count_columns = ['Category'] + [f'Cable_{val}' for val in sorted(unique_values)]
    counts_df = pd.DataFrame(counts_list, columns=count_columns)

    # Write both DataFrames to the same Excel file in the same sheet
    excel_filename = os.path.join(hh_run_path, 'rescaled_positions.xlsx')
    with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
        positions_df.to_excel(writer, sheet_name='Positions_and_Counts', index=False)

        # Start writing counts DataFrame below the positions DataFrame
        startrow = len(positions_df) + 2  # +2 for a blank line between tables
        counts_df.to_excel(writer, sheet_name='Positions_and_Counts', startrow=startrow, index=False)

    # Apply conditional formatting to the Excel file.
    format_configuration_sheet(excel_filename, positions_df, nr_of_backbone_switches)